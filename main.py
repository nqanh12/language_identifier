# main.py
import tkinter as tk
from tkinter import messagebox, ttk, filedialog
import joblib
import json
import webbrowser
from pathlib import Path
import matplotlib.pyplot as plt
from matplotlib.backends.backend_tkagg import FigureCanvasTkAgg
import csv
from datetime import datetime
import markdown2
from reportlab.pdfgen import canvas
from reportlab.lib.pagesizes import letter
from reportlab.pdfbase import pdfmetrics
from reportlab.pdfbase.ttfonts import TTFont
import subprocess
import sys
from ttkthemes import ThemedTk
import sv_ttk
from PIL import Image, ImageTk
import darkdetect
try:
    import docx
except ImportError:
    docx = None
try:
    import PyPDF2
except ImportError:
    PyPDF2 = None

# Load model and vectorizer
model = joblib.load('model/nb_model.pkl')
vectorizer = joblib.load('model/vectorizer.pkl')

# Load feature importance
with open('model/feature_importance.json', 'r', encoding='utf-8') as f:
    feature_importance = json.load(f)

# Initialize history file if it doesn't exist
history_file = 'history.csv'
if not Path(history_file).exists():
    with open(history_file, 'w', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow(['Timestamp', 'Text', 'Predicted Language', 'Probability'])

def export_to_pdf():
    """Export the model report to PDF format"""
    report_path = Path('reports/model_report.md')
    if not report_path.exists():
        messagebox.showwarning("Cảnh báo", "Không tìm thấy file báo cáo!")
        return
        
    # Read markdown content
    with open(report_path, 'r', encoding='utf-8') as f:
        md_content = f.read()
    
    # Convert markdown to HTML
    html_content = markdown2.markdown(md_content)
    
    # Create PDF
    pdf_path = 'reports/model_report.pdf'
    c = canvas.Canvas(pdf_path, pagesize=letter)
    width, height = letter
    
    # Add title
    c.setFont("Helvetica-Bold", 14)
    c.drawString(50, height - 50, "Language Detection System Report")
    
    # Add content
    c.setFont("Helvetica", 10)
    y = height - 100
    for line in html_content.split('\n'):
        if line.strip():
            if line.startswith('<h1>'):
                c.setFont("Helvetica-Bold", 14)
                text = line.replace('<h1>', '').replace('</h1>', '')
                c.drawString(50, y, text)
                y -= 25
            elif line.startswith('<h2>'):
                c.setFont("Helvetica-Bold", 12)
                text = line.replace('<h2>', '').replace('</h2>', '')
                c.drawString(50, y, text)
                y -= 20
            elif line.startswith('<h3>'):
                c.setFont("Helvetica-Bold", 10)
                text = line.replace('<h3>', '').replace('</h3>', '')
                c.drawString(50, y, text)
                y -= 15
            else:
                c.setFont("Helvetica", 10)
                text = line.replace('<p>', '').replace('</p>', '')
                c.drawString(50, y, text)
                y -= 12
            
            # Check if we need a new page
            if y < 50:
                c.showPage()
                y = height - 50
    
    c.save()
    messagebox.showinfo("Thông báo", f"Đã xuất báo cáo PDF thành công: {pdf_path}")

def show_history():
    """Show prediction history in a new window"""
    history_win = tk.Toplevel(root)
    history_win.title("Lịch sử dự đoán")
    history_win.geometry("600x400")
    
    # Create treeview
    tree = ttk.Treeview(history_win, columns=("Thời gian", "Văn bản", "Ngôn ngữ", "Xác suất"), show='headings')
    tree.heading("Thời gian", text="Thời gian")
    tree.heading("Văn bản", text="Văn bản")
    tree.heading("Ngôn ngữ", text="Ngôn ngữ")
    tree.heading("Xác suất", text="Xác suất")
    
    # Set column widths
    tree.column("Thời gian", width=120)
    tree.column("Văn bản", width=300)
    tree.column("Ngôn ngữ", width=80)
    tree.column("Xác suất", width=80)
    
    # Add scrollbar
    scrollbar = ttk.Scrollbar(history_win, orient="vertical", command=tree.yview)
    tree.configure(yscrollcommand=scrollbar.set)
    
    # Pack widgets
    tree.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
    scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
    
    # Load history
    try:
        with open(history_file, 'r', newline='', encoding='utf-8') as f:
            reader = csv.reader(f)
            next(reader)  # Skip header
            for row in reader:
                tree.insert('', tk.END, values=row)
    except Exception as e:
        messagebox.showerror("Lỗi", f"Không thể đọc lịch sử: {str(e)}")

def predict_language():
    # Lấy text từ input
    text = input_text.get("1.0", tk.END).strip()
    if not text:
        messagebox.showwarning("Cảnh báo", "Vui lòng nhập văn bản!")
        return
    
    # Chuyển đổi text thành vector features
    text_vec = vectorizer.transform([text])
    
    # Dự đoán và tính xác suất
    prediction = model.predict(text_vec)[0]
    probabilities = model.predict_proba(text_vec)[0]
    max_prob = probabilities.max() * 100
    
    # Cập nhật kết quả lên GUI
    result_label.config(text=f"Ngôn ngữ: {prediction} ({max_prob:.2f}%)")
    
    # Cập nhật thanh xác suất
    for i, (lang, prob) in enumerate(zip(model.classes_, probabilities)):
        prob_labels[i].config(text=f"{lang}: {prob*100:.2f}%")
        prob_bars[i]['value'] = prob * 100
    
    # Save to history
    with open(history_file, 'a', newline='', encoding='utf-8') as f:
        writer = csv.writer(f)
        writer.writerow([datetime.now().strftime("%Y-%m-%d %H:%M:%S"), 
                        text[:100] + "..." if len(text) > 100 else text,
                        prediction,
                        f"{max_prob:.2f}%"])
    
    # Show top features
    show_top_features_for_input(text, prediction)

def show_feature_importance():
    # Create new window
    feature_window = tk.Toplevel(root)
    feature_window.title("Từ đặc trưng cho mỗi ngôn ngữ")
    feature_window.geometry("600x400")
    
    # Create notebook for tabs
    notebook = ttk.Notebook(feature_window)
    notebook.pack(fill='both', expand=True, padx=5, pady=3)
    
    # Create tab for each language
    for lang in model.classes_:
        frame = ttk.Frame(notebook)
        notebook.add(frame, text=lang)
        
        # Create text widget with scrollbar
        text_widget = tk.Text(frame, wrap=tk.WORD, padx=5, pady=3, font=("Arial", 10))
        scrollbar = ttk.Scrollbar(frame, orient="vertical", command=text_widget.yview)
        text_widget.configure(yscrollcommand=scrollbar.set)
        
        # Pack widgets
        text_widget.pack(side=tk.LEFT, fill=tk.BOTH, expand=True)
        scrollbar.pack(side=tk.RIGHT, fill=tk.Y)
        
        # Add features
        text_widget.insert(tk.END, f"Top 20 từ đặc trưng cho {lang}:\n\n")
        for feat in feature_importance[lang]:
            text_widget.insert(tk.END, f"- {feat['feature']}: {feat['importance']:.2f}\n")

def show_report():
    report_path = Path('reports/model_report.md')
    if report_path.exists():
        webbrowser.open(report_path)
    else:
        messagebox.showwarning("Cảnh báo", "Không tìm thấy file báo cáo!")

def show_top_features_for_input(text, predicted_lang):
    # Lấy các n-gram/từ xuất hiện trong văn bản đầu vào
    features = vectorizer.transform([text])
    feature_names = vectorizer.get_feature_names_out()
    input_indices = features.nonzero()[1]
    # Lấy trọng số của các feature này với ngôn ngữ dự đoán
    lang_idx = list(model.classes_).index(predicted_lang)
    feature_log_prob = model.feature_log_prob_[lang_idx]
    feature_scores = []
    for idx in input_indices:
        feature_scores.append((feature_names[idx], feature_log_prob[idx]))
    # Sắp xếp theo trọng số giảm dần
    feature_scores.sort(key=lambda x: x[1], reverse=True)
    # Lấy top 10
    top_features = feature_scores[:10]
    # Hiển thị trong cửa sổ mới
    top_window = tk.Toplevel(root)
    top_window.title(f"Top n-gram ảnh hưởng nhất ({predicted_lang})")
    top_window.geometry("300x250")
    label = tk.Label(top_window, text=f"Top n-gram/từ ảnh hưởng nhất cho {predicted_lang}:", font=("Arial", 10, "bold"))
    label.pack(pady=5)
    text_widget = tk.Text(top_window, height=12, width=35, font=("Arial", 10))
    for feat, score in top_features:
        text_widget.insert(tk.END, f"- {feat}: {score:.2f}\n")
    text_widget.pack(padx=5, pady=3)
    text_widget.config(state=tk.DISABLED)

def show_probability_chart():
    # Lấy xác suất hiện tại
    text = input_text.get("1.0", tk.END).strip()
    if not text:
        messagebox.showwarning("Cảnh báo", "Vui lòng nhập văn bản!")
        return
    text_vec = vectorizer.transform([text])
    probabilities = model.predict_proba(text_vec)[0]
    langs = model.classes_
    # Vẽ biểu đồ
    fig, ax = plt.subplots(figsize=(5, 3))
    bars = ax.bar(langs, probabilities * 100, color='#3498db')
    ax.set_ylabel('Xác suất (%)')
    ax.set_title('Xác suất dự đoán cho từng ngôn ngữ')
    ax.set_ylim(0, 100)
    for bar, prob in zip(bars, probabilities):
        ax.text(bar.get_x() + bar.get_width()/2, bar.get_height(), f'{prob*100:.2f}%', ha='center', va='bottom', fontsize=8)
    # Hiển thị trong cửa sổ Tkinter
    chart_window = tk.Toplevel(root)
    chart_window.title("Biểu đồ xác suất ngôn ngữ")
    chart_window.geometry("600x300")
    canvas = FigureCanvasTkAgg(fig, master=chart_window)
    canvas.draw()
    canvas.get_tk_widget().pack(fill=tk.BOTH, expand=True)
    plt.close(fig)

# Xóa văn bản và kết quả
def clear_text():
    input_text.delete("1.0", tk.END)
    result_label.config(text="")
    for label in prob_labels:
        label.config(text="")
    for bar in prob_bars:
        bar['value'] = 0

# Đọc nội dung file
def read_file_content(filepath):
    ext = filepath.lower().split('.')[-1]
    if ext == 'txt':
        with open(filepath, 'r', encoding='utf-8') as f:
            return f.read()
    elif ext == 'docx' and docx:
        doc = docx.Document(filepath)
        return '\n'.join([p.text for p in doc.paragraphs])
    elif ext == 'pdf' and PyPDF2:
        text = ''
        with open(filepath, 'rb') as f:
            reader = PyPDF2.PdfReader(f)
            for page in reader.pages:
                text += page.extract_text() + '\n'
        return text
    else:
        messagebox.showerror("Lỗi", f"Không hỗ trợ định dạng file: {ext} hoặc thiếu thư viện!")
        return None

# Nhận diện từ file
def detect_from_file():
    filepath = filedialog.askopenfilename(filetypes=[
        ("Text files", "*.txt"),
        ("Word files", "*.docx"),
        ("PDF files", "*.pdf"),
        ("All files", "*.*")
    ])
    if not filepath:
        return
    content = read_file_content(filepath)
    if not content:
        return
    # Dự đoán
    text_vec = vectorizer.transform([content])
    prediction = model.predict(text_vec)[0]
    probabilities = model.predict_proba(text_vec)[0]
    # Hiển thị kết quả
    result = f"File: {filepath}\nNgôn ngữ: {prediction} ({probabilities.max()*100:.2f}%)\n"
    result += "\nXác suất:\n"
    for lang, prob in zip(model.classes_, probabilities):
        result += f"- {lang}: {prob*100:.2f}%\n"
    # Hiển thị cửa sổ kết quả
    file_result_win = tk.Toplevel(root)
    file_result_win.title("Kết quả nhận diện từ file")
    file_result_win.geometry("400x300")
    text_widget = tk.Text(file_result_win, font=("Arial", 10))
    text_widget.insert(tk.END, result)
    text_widget.pack(fill=tk.BOTH, expand=True)
    text_widget.config(state=tk.DISABLED)

# Batch nhận diện nhiều đoạn văn
def batch_detect():
    def do_batch_detect():
        input_data = batch_text.get("1.0", tk.END).strip().splitlines()
        if not input_data:
            messagebox.showwarning("Cảnh báo", "Vui lòng nhập các đoạn văn!")
            return
        results = []
        for line in input_data:
            if not line.strip():
                results.append(("", "", ""))
                continue
            text_vec = vectorizer.transform([line])
            prediction = model.predict(text_vec)[0]
            probabilities = model.predict_proba(text_vec)[0]
            prob = probabilities[list(model.classes_).index(prediction)]
            results.append((line, prediction, f"{prob*100:.2f}%"))
        # Hiển thị kết quả dạng bảng
        result_win = tk.Toplevel(batch_win)
        result_win.title("Kết quả batch")
        result_win.geometry("600x300")
        tree = ttk.Treeview(result_win, columns=("Văn bản", "Ngôn ngữ", "Xác suất"), show='headings')
        tree.heading("Văn bản", text="Văn bản")
        tree.heading("Ngôn ngữ", text="Ngôn ngữ")
        tree.heading("Xác suất", text="Xác suất")
        tree.column("Văn bản", width=300)
        tree.column("Ngôn ngữ", width=100)
        tree.column("Xác suất", width=80)
        for row in results:
            tree.insert('', tk.END, values=row)
        tree.pack(fill=tk.BOTH, expand=True)

    batch_win = tk.Toplevel(root)
    batch_win.title("Batch Nhận Diện Ngôn Ngữ")
    batch_win.geometry("500x300")
    label = tk.Label(batch_win, text="Nhập mỗi đoạn văn trên một dòng:", font=("Arial", 10))
    label.pack(pady=5)
    batch_text = tk.Text(batch_win, height=8, font=("Arial", 10))
    batch_text.pack(fill=tk.BOTH, expand=True, padx=5)
    detect_btn = ttk.Button(batch_win, text="Nhận Diện", command=do_batch_detect, style="Custom.TButton")
    detect_btn.pack(pady=5)

def add_training_data():
    """Open window for adding new training data"""
    train_win = tk.Toplevel(root)
    train_win.title("Thêm dữ liệu huấn luyện")
    train_win.geometry("400x350")
    
    # Language selection
    lang_frame = ttk.Frame(train_win)
    lang_frame.pack(fill=tk.X, padx=5, pady=3)
    ttk.Label(lang_frame, text="Chọn ngôn ngữ:").pack(side=tk.LEFT)
    lang_var = tk.StringVar()
    lang_combo = ttk.Combobox(lang_frame, textvariable=lang_var, values=list(model.classes_))
    lang_combo.pack(side=tk.LEFT, padx=3)
    lang_combo.set(model.classes_[0])
    
    # Text input
    text_frame = ttk.LabelFrame(train_win, text="Nhập câu mẫu", padding="5")
    text_frame.pack(fill=tk.BOTH, expand=True, padx=5, pady=3)
    
    text_input = tk.Text(text_frame, height=8, font=("Arial", 10), fg="black")
    text_input.pack(fill=tk.BOTH, expand=True)
    
    def save_and_train():
        lang = lang_var.get()
        text = text_input.get("1.0", tk.END).strip()
        if not text:
            messagebox.showwarning("Cảnh báo", "Vui lòng nhập câu mẫu!")
            return
            
        # Save to data file
        data_file = f'data/{lang.lower()[:2]}.txt'
        with open(data_file, 'a', encoding='utf-8') as f:
            f.write(text + '\n')
            
        # Retrain model
        try:
            subprocess.run([sys.executable, 'train_model.py'], check=True)
            messagebox.showinfo("Thông báo", "Đã thêm dữ liệu và huấn luyện lại mô hình thành công!")
            train_win.destroy()
            
            # Reload model and vectorizer
            global model, vectorizer
            model = joblib.load('model/nb_model.pkl')
            vectorizer = joblib.load('model/vectorizer.pkl')
            
            # Reload feature importance
            global feature_importance
            with open('model/feature_importance.json', 'r', encoding='utf-8') as f:
                feature_importance = json.load(f)
                
        except subprocess.CalledProcessError as e:
            messagebox.showerror("Lỗi", f"Không thể huấn luyện lại mô hình: {str(e)}")
    
    # Save button
    save_btn = ttk.Button(
        train_win,
        text="Lưu và Huấn luyện",
        command=save_and_train,
        style="Custom.TButton"
    )
    save_btn.pack(pady=5)

# Giao diện Tkinter
root = ThemedTk(theme="arc")  # Use a modern theme
root.title("Language Detector - Naive Bayes")
root.geometry("1400x800")

# Set dark/light theme based on system
if darkdetect.isDark():
    sv_ttk.set_theme("dark")
else:
    sv_ttk.set_theme("light")

# Get screen dimensions
screen_width = root.winfo_screenwidth()
screen_height = root.winfo_screenheight()

# Calculate position coordinates
x = (screen_width - 1400) // 2
y = (screen_height - 800) // 2

# Set window position
root.geometry(f"1400x800+{x}+{y}")

# Configure styles
style = ttk.Style()
style.configure("Custom.TButton", 
                padding=10, 
                font=("Segoe UI", 10),
                background="#3498db",
                foreground="black")

style.configure("Title.TLabel",
                font=("Segoe UI", 28, "bold"),
                foreground="black",
                padding=15)

style.configure("Subtitle.TLabel",
                font=("Segoe UI", 11),
                foreground="black",
                padding=8)

style.configure("Result.TLabel",
                font=("Segoe UI", 16, "bold"),
                foreground="black",
                padding=8)

style.configure("Custom.TLabelframe",
                font=("Segoe UI", 11, "bold"),
                padding=12)

style.configure("Custom.TLabelframe.Label",
                font=("Segoe UI", 11, "bold"),
                foreground="black")

style.configure("Custom.Horizontal.TProgressbar",
                troughcolor='#f0f0f0',
                background='#3498db',
                thickness=25)

# Main container with modern padding
main_container = ttk.Frame(root, padding="25")
main_container.pack(fill=tk.BOTH, expand=True)

# Title with modern styling
title_frame = ttk.Frame(main_container)
title_frame.pack(fill=tk.X, pady=(0, 25))

title_label = ttk.Label(
    title_frame,
    text="Language Detection System",
    style="Title.TLabel"
)
title_label.pack(side=tk.LEFT)

# Theme toggle button
def toggle_theme():
    if sv_ttk.get_theme() == "dark":
        sv_ttk.set_theme("light")
    else:
        sv_ttk.set_theme("dark")

theme_btn = ttk.Button(
    title_frame,
    text="🌓 Toggle Theme",
    command=toggle_theme,
    style="Custom.TButton"
)
theme_btn.pack(side=tk.RIGHT)

# Create two main columns with modern spacing
content_frame = ttk.Frame(main_container)
content_frame.pack(fill=tk.BOTH, expand=True)

# Left column (Input and Controls)
left_column = ttk.Frame(content_frame, padding="15")
left_column.pack(side=tk.LEFT, fill=tk.BOTH, expand=True, padx=(0, 15))

# Input frame with modern styling
input_frame = ttk.LabelFrame(left_column, text="Input Text", style="Custom.TLabelframe")
input_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 15))

input_text = tk.Text(
    input_frame,
    height=8,
    font=("Segoe UI", 12),
    wrap=tk.WORD,
    bd=2,
    relief=tk.SOLID,
    padx=15,
    pady=15,
    fg="black"
)
input_text.pack(fill=tk.BOTH, expand=True, padx=10, pady=10)

# Add placeholder text
def on_focus_in(event):
    if input_text.get("1.0", tk.END).strip() == "Enter text here...":
        input_text.delete("1.0", tk.END)
        input_text.config(fg="black")

def on_focus_out(event):
    if not input_text.get("1.0", tk.END).strip():
        input_text.insert("1.0", "Enter text here...")
        input_text.config(fg="gray")

input_text.insert("1.0", "Enter text here...")
input_text.config(fg="gray")
input_text.bind("<FocusIn>", on_focus_in)
input_text.bind("<FocusOut>", on_focus_out)

# Control buttons frame with modern layout
control_frame = ttk.Frame(left_column)
control_frame.pack(fill=tk.X, pady=15)

# Primary buttons with modern styling
primary_buttons = ttk.Frame(control_frame)
primary_buttons.pack(fill=tk.X, pady=(0, 10))

predict_btn = ttk.Button(
    primary_buttons,
    text="🔍 Detect Language",
    command=predict_language,
    style="Custom.TButton"
)
predict_btn.pack(side=tk.LEFT, padx=3)

clear_btn = ttk.Button(
    primary_buttons,
    text="🗑️ Clear",
    command=clear_text,
    style="Custom.TButton"
)
clear_btn.pack(side=tk.LEFT, padx=3)

# Secondary buttons with modern styling
secondary_buttons = ttk.Frame(control_frame)
secondary_buttons.pack(fill=tk.X)

file_btn = ttk.Button(
    secondary_buttons,
    text="📄 Detect from File",
    command=detect_from_file,
    style="Custom.TButton"
)
file_btn.pack(side=tk.LEFT, padx=3)

batch_btn = ttk.Button(
    secondary_buttons,
    text="📋 Batch Detection",
    command=batch_detect,
    style="Custom.TButton"
)
batch_btn.pack(side=tk.LEFT, padx=3)

# Right column (Results and Analysis)
right_column = ttk.Frame(content_frame, padding="15")
right_column.pack(side=tk.RIGHT, fill=tk.BOTH, expand=True, padx=(15, 0))

# Result frame with modern styling
result_frame = ttk.LabelFrame(right_column, text="Detection Result", style="Custom.TLabelframe")
result_frame.pack(fill=tk.X, pady=(0, 15))

result_label = ttk.Label(
    result_frame,
    text="",
    style="Result.TLabel"
)
result_label.pack(pady=15)

# Probability frame with modern styling
prob_frame = ttk.LabelFrame(right_column, text="Language Probabilities", style="Custom.TLabelframe")
prob_frame.pack(fill=tk.BOTH, expand=True, pady=(0, 15))

prob_labels = []
prob_bars = []

for lang in model.classes_:
    frame = ttk.Frame(prob_frame)
    frame.pack(fill=tk.X, pady=5, padx=10)
    
    label = ttk.Label(frame, text="", width=15, font=("Segoe UI", 11))
    label.pack(side=tk.LEFT, padx=5)
    prob_labels.append(label)
    
    bar = ttk.Progressbar(frame, 
                         length=300,
                         mode='determinate',
                         style="Custom.Horizontal.TProgressbar")
    bar.pack(side=tk.LEFT, padx=5, fill=tk.X, expand=True)
    prob_bars.append(bar)

# Analysis buttons frame with modern styling
analysis_frame = ttk.Frame(right_column)
analysis_frame.pack(fill=tk.X, pady=15)

features_btn = ttk.Button(
    analysis_frame,
    text="📊 View Features",
    command=show_feature_importance,
    style="Custom.TButton"
)
features_btn.pack(side=tk.LEFT, padx=3)

prob_chart_btn = ttk.Button(
    analysis_frame,
    text="📈 Probability Chart",
    command=show_probability_chart,
    style="Custom.TButton"
)
prob_chart_btn.pack(side=tk.LEFT, padx=3)

# Additional tools frame with modern styling
tools_frame = ttk.Frame(right_column)
tools_frame.pack(fill=tk.X, pady=15)

report_btn = ttk.Button(
    tools_frame,
    text="📑 View Report",
    command=show_report,
    style="Custom.TButton"
)
report_btn.pack(side=tk.LEFT, padx=3)

export_pdf_btn = ttk.Button(
    tools_frame,
    text="📥 Export PDF",
    command=export_to_pdf,
    style="Custom.TButton"
)
export_pdf_btn.pack(side=tk.LEFT, padx=3)

history_btn = ttk.Button(
    tools_frame,
    text="📜 View History",
    command=show_history,
    style="Custom.TButton"
)
history_btn.pack(side=tk.LEFT, padx=3)

train_btn = ttk.Button(
    tools_frame,
    text="➕ Add Training Data",
    command=add_training_data,
    style="Custom.TButton"
)
train_btn.pack(side=tk.LEFT, padx=3)

# Info label at bottom with modern styling
info_frame = ttk.Frame(main_container)
info_frame.pack(fill=tk.X, pady=15)

info_label = ttk.Label(
    info_frame,
    text="Supported Languages: Vietnamese, English, French, Japanese, German, Spanish, Korean",
    style="Subtitle.TLabel"
)
info_label.pack()

# Add tooltips
def create_tooltip(widget, text):
    def show_tooltip(event):
        tooltip = tk.Toplevel()
        tooltip.wm_overrideredirect(True)
        tooltip.wm_geometry(f"+{event.x_root+10}+{event.y_root+10}")
        
        label = ttk.Label(tooltip, text=text, justify=tk.LEFT,
                         background="#ffffe0", relief=tk.SOLID, borderwidth=1)
        label.pack()
        
        def hide_tooltip():
            tooltip.destroy()
            
        widget.bind('<Leave>', lambda e: hide_tooltip())
        tooltip.bind('<Leave>', lambda e: hide_tooltip())
        
    widget.bind('<Enter>', show_tooltip)

# Add tooltips to buttons
create_tooltip(predict_btn, "Detect the language of the input text")
create_tooltip(clear_btn, "Clear the input text and results")
create_tooltip(file_btn, "Detect language from a text file")
create_tooltip(batch_btn, "Detect language for multiple text entries")
create_tooltip(features_btn, "View the most important features for each language")
create_tooltip(prob_chart_btn, "View a chart of language probabilities")
create_tooltip(report_btn, "View the model performance report")
create_tooltip(export_pdf_btn, "Export the report to PDF format")
create_tooltip(history_btn, "View prediction history")
create_tooltip(train_btn, "Add new training data to improve the model")
create_tooltip(theme_btn, "Toggle between light and dark theme")

root.mainloop()
